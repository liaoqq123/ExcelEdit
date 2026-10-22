"""结果导出的原文保留与原子写入测试。"""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from result_export import export_result_rows


class ResultExportTests(unittest.TestCase):
    def test_export_preserves_literal_formulas_unicode_and_multiline_text(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "检索结果.xlsx"
            values = ["=1+1", "+SUM(A1:A2)", "-123", "@命令", "中文\n第二行\t内容", 42, True]
            count = export_result_rows(output, ["=原始内容", "内容"], (values for _ in range(2)))

            self.assertEqual(count, 2)
            workbook = load_workbook(output, read_only=True, data_only=False)
            try:
                rows = list(workbook.active.iter_rows())
                self.assertEqual(rows[0][0].value, "=原始内容")
                self.assertEqual(rows[0][0].data_type, "s")
                self.assertEqual([cell.value for cell in rows[1]], [str(value) for value in values])
                self.assertTrue(all(cell.data_type == "s" for cell in rows[1]))
                self.assertEqual([cell.value for cell in rows[2]], [str(value) for value in values])
            finally:
                workbook.close()
            self.assertEqual(list(Path(directory).iterdir()), [output])

    def test_iterator_failure_keeps_existing_file_and_cleans_temporary_output(self) -> None:
        def failing_rows():
            yield ["第一条"]
            raise RuntimeError("读取结果失败")

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "结果.xlsx"
            original = b"existing workbook content"
            output.write_bytes(original)
            previous_internal_files = set(ALL_TEMP_FILES)
            with self.assertRaisesRegex(RuntimeError, "读取结果失败"):
                export_result_rows(output, ["内容"], failing_rows())

            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(list(Path(directory).iterdir()), [output])
            self.assertEqual(set(ALL_TEMP_FILES), previous_internal_files)

    def test_replace_failure_keeps_existing_file_and_cleans_temporary_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "结果.xlsx"
            original = b"existing workbook content"
            output.write_bytes(original)
            with patch("result_export.os.replace", side_effect=PermissionError("文件被占用")):
                with self.assertRaises(PermissionError):
                    export_result_rows(output, ["内容"], [["新内容"]])

            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(list(Path(directory).iterdir()), [output])

    def test_overlong_cell_fails_without_silent_truncation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "结果.xlsx"
            with self.assertRaisesRegex(ValueError, "32767"):
                export_result_rows(output, ["内容"], [["长" * 32768]])
            self.assertEqual(list(Path(directory).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
